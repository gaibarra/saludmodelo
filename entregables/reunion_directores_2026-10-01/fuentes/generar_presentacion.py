from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
from html import escape
import base64
from PIL import Image
O=Path('/home/gaibarra/plandetrabajo/entregables/reunion_directores_2026-10-01');SRC=O.parent/'reunion_directores_2026-09-30'
r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
BG='F3F9F8';INK='173F49';TEAL='087F82';MUT='54717B';WHITE='FFFFFF';PALE='E2F1ED';GOLD='FFF1CB'
slides=[];notes=[];s=None;parts=[]
def rect(x,y,w,h,color):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
 parts.append(f'<div style="position:absolute;left:{x}in;top:{y}in;width:{w}in;height:{h}in;background:#{color}"></div>')
def txt(x,y,w,h,text,size=22,color=INK,bold=False):
 sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(8)
 parts.append(f'<div style="position:absolute;left:{x}in;top:{y}in;width:{w}in;height:{h}in;font-size:{size}pt;color:#{color};font-weight:{700 if bold else 400};line-height:1.18;white-space:pre-line">{escape(text)}</div>')
def new(title,kicker,note):
 global s,parts
 if s is not None:slides.append(''.join(parts))
 s=r.slides.add_slide(r.slide_layouts[6]);parts=[];s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 rect(0,0,.16,7.5,TEAL);txt(.55,.32,12,.3,'UNIVERSIDAD MODELO  /  SALUD Y ODONTOLOGÍA',10,TEAL,True)
 txt(.55,.87,12.1,1,title,30,INK,True);txt(.58,1.9,12,.5,kicker,16,MUT)
 txt(.58,7.12,11,.2,'REUNIÓN CON DIRECCIONES · PROPUESTA DE TRABAJO · 01 OCT 2026',9,MUT)
 txt(12.15,7.1,.6,.25,f'{len(r.slides):02d}',11,TEAL,True)
 s.notes_slide.notes_text_frame.text=note;notes.append(f'{len(r.slides):02d}. {title}\n{note}\n')
def card(x,y,w,title,body,h=2.3):
 rect(x,y,w,h,WHITE);rect(x,y,.055,h,TEAL);txt(x+.2,y+.2,w-.4,.6,title,21,TEAL,True);txt(x+.2,y+(.68 if h<1.8 else .98),w-.4,h-(.74 if h<1.8 else 1.05),body,17 if h<1.8 else 18)
def banner(text,y=6.3):rect(.58,y,12.1,.57,PALE);txt(.78,y+.12,11.7,.4,text,14,TEAL,True)
def photo(name,x,y,w,h):
 path=(O/'imagenes'/name) if (O/'imagenes'/name).exists() else SRC/name;im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale
 x+=(w-nw)/2;y+=(h-nh)/2;s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(nw),height=Inches(nh))
 b=base64.b64encode(path.read_bytes()).decode();parts.append(f'<img src="data:image/png;base64,{b}" style="position:absolute;left:{x}in;top:{y}in;width:{nw}in;height:{nh}in">')
new('¿Qué es la aplicación Salud y Odontología?','Una plataforma web para organizar servicios y acompañar la formación de los alumnos.','Presentar la aplicación como una plataforma institucional que conecta la atención de los servicios, la organización del personal y la formación académica. Participan pacientes y usuarios, colaboradores, alumnos, supervisores y Direcciones, cada uno con permisos específicos. El alcance previsto incluye solicitudes de atención, seguimiento y expedientes por servicio, vinculación de prácticas con competencias e informes individuales y consolidados. Las escuelas conservan su administración independiente. Aclarar que el despliegue es por etapas: la demostración y el núcleo académico sirven de base, mientras que la operación clínica real requiere validación y completar los módulos correspondientes.')
txt(.7,2.55,11.9,1.22,'Conecta a pacientes y usuarios, colaboradores, alumnos, supervisores y Direcciones para dar seguimiento al servicio prestado y al aprendizaje que genera.',23,INK)
card(.65,3.85,3.9,'Atención y servicios','Solicitudes de atención, seguimiento y expedientes adaptados a cada servicio.',2.15)
card(4.72,3.85,3.9,'Formación','Prácticas, supervisión, competencias e informes por alumno y ciclo.',2.15)
card(8.79,3.85,3.9,'Gestión y Dirección','Responsables, tareas y avances; cada Escuela con administración propia.',2.15)
banner('Implementación por etapas: demostración y núcleo académico; alcance clínico por validar y completar.')
new('El objetivo: conocer y respaldar cada práctica','Una visión compartida para las Direcciones de Salud y Odontología.','Abrir con el objetivo académico. Preguntar qué información resulta indispensable para cada Dirección durante el ciclo y al finalizarlo. No presentar las cifras simuladas como resultados reales.')
txt(.7,2.7,11.7,1.7,'Saber qué hizo cada alumno,\nquién lo supervisó y qué competencias alcanzó.',32,INK,True)
card(.7,4.7,5.8,'Durante el ciclo','Avances, pendientes y seguimiento oportuno.',1.35)
card(6.7,4.7,5.8,'Al finalizar el ciclo','Informes individuales y consolidados.',1.35)
banner('Hoy buscamos acuerdos para comenzar con personas y procesos reales.')
new('Lo que Dirección necesita poder responder','Primero acordamos la información útil; después afinamos los formularios.','Pedir a los Directores que prioricen estas cuatro preguntas y agreguen lo que falte. Acordar qué información será necesaria para dar por cerrado el ciclo académico.')
card(.65,2.7,5.9,'01 · Participación','¿Qué prácticas realizó cada alumno y en qué servicio?',1.5)
card(6.75,2.7,5.9,'02 · Supervisión','¿Quién revisó y validó cada actividad?',1.5)
card(.65,4.45,5.9,'03 · Aprendizaje','¿Qué competencias demostró y cuáles debe reforzar?',1.5)
card(6.75,4.45,5.9,'04 · Cierre','¿Qué resultados y pendientes deja el ciclo?',1.5)
new('Una aplicación que conecta el trabajo cotidiano','Del registro de una actividad al informe académico.','Explicar el flujo en lenguaje sencillo. Una captura no equivale automáticamente a una práctica acreditada: debe existir revisión del supervisor. Los módulos académicos sirven de base; los formatos clínicos siguen sujetos a validación y desarrollo.')
for x,title,body in [( .65,'1 · Registrar','Actividad, alumno, servicio y fecha.'),(3.75,'2 · Supervisar','Revisión, observaciones y correcciones.'),(6.85,'3 · Validar','Actividad y competencia acreditadas.'),(9.95,'4 · Consultar','Avances e informes para Dirección.')]:card(x,2.9,2.75,title,body,2.6)
banner('La validación del supervisor convierte el registro en evidencia académica.')
new('Dos escuelas, administración independiente','Salud y Odontología tienen el mismo nivel jerárquico.','Confirmar la separación administrativa. Cada Dirección organiza su escuela, sus responsables y el seguimiento de sus alumnos.')
card(.7,2.8,5.8,'Escuela de Salud','Su Dirección administra sus servicios, colaboradores y seguimiento académico.',2.3)
card(6.75,2.8,5.8,'Escuela de Odontología','Su Dirección administra su equipo, operación y prácticas odontológicas.',2.3)
txt(.9,5.5,11.5,.7,'Cada Escuela define sus responsables y administra su información,\ncon funciones y permisos propios.',20)
new('Así se podrá consultar el avance · Salud','Vista ilustrativa de la interfaz actual con cifras ficticias.','Todas las cifras de esta imagen son simuladas y corresponden únicamente al ejemplo de la Escuela de Salud. No se cargaron registros ficticios en producción para estas imágenes.')
photo('02_Direccion_Salud.png',.6,2.48,8.5,4.38)
txt(9.3,2.9,3.3,3.4,'Administración de Salud.\n\nPrácticas supervisadas.\n\nAvance de sus alumnos.',21,INK,True)
new('Así se podrá consultar el avance · Odontología','Vista ilustrativa de la interfaz actual con cifras ficticias.','Se muestra la Escuela de Odontología con administración propia. Las cifras ilustran el funcionamiento posterior, no los registros reales ni una meta institucional. Los paneles completos también se entregan como PNG y PDF aparte.')
photo('03_Direccion_Odontologia.png',.6,2.48,8.5,4.38)
txt(9.3,2.9,3.3,3.4,'Administración propia.\n\nPrácticas supervisadas.\n\nInformes del ciclo.',21,INK,True)
new('Octubre: colaboradores y roles completos','Del 1 al 31 de octubre · Cada Dirección elige aplicación o Excel.','Extender la recopilación y revisión a todo octubre, con meta de cierre al 31 de octubre. Cada Dirección decide si su equipo captura en la aplicación o trabaja en la plantilla Excel. Identificar colaboradores, servicios, roles, participación y vigencia; revisar responsables, duplicados y personas que participan en más de un servicio. Consolidar y validar la relación con cada Dirección. Excel es una alternativa de recopilación: la carga debe prepararse y validarse, no se presenta como una importación automática ya disponible. No pedir contraseñas ni información clínica en Excel.')
card(.65,2.75,3.9,'Qué registrar','Nombre y correo institucional.\nServicio y rol asignado.\nParticipación y vigencia.',2.8)
card(4.72,2.75,3.9,'Cómo colaborar','Aplicación o plantilla Excel, a elección de cada Dirección.\nUn enlace por Escuela.',2.8)
card(8.79,2.75,3.9,'Cómo cerrar','Revisar y completar durante octubre.\nValidar colaboradores y roles al 31 de octubre.',2.8)
banner('Un mismo plazo: todo octubre para recopilar, corregir y validar la información.')
new('Mi compromiso: visitas durante todo octubre','C.P. Gonzalo Arturo Ibarra Mendoza · Trabajo presencial, completo y a conciencia.','Distribuir las visitas presenciales durante todo el mes de octubre. Se mantiene el compromiso de al menos dos días por servicio, ampliando visitas y seguimiento cuando sea necesario. El objetivo es conocer los procesos a conciencia y completar el recorrido de todos los servicios, procurando no dejar asuntos pendientes. Escuchar al equipo, observar el trabajo, revisar excepciones y confirmar lo comprendido. Coordinar horarios sin interferir con la atención. Cualquier pendiente que no pueda resolverse al cierre de octubre quedará documentado con responsable y fecha de seguimiento.')
card(.7,2.8,5.8,'Conocer y comprender','Observar el proceso actual.\nEscuchar al equipo y sus necesidades.\nRevisar formatos y excepciones.',2.65)
card(6.75,2.8,5.8,'Validar y completar','Confirmar lo observado con el servicio.\nVolver cuando sea necesario.\nResolver dudas y acordar ajustes.',2.65)
banner('Al menos dos días por servicio, con seguimiento durante octubre para procurar un trabajo completo.')
new('Octubre: visitas y registro en paralelo','Plazo común: del 1 al 31 de octubre de 2026 · Agenda por acordar con cada servicio.','Usar todo octubre tanto para las visitas como para completar colaboradores y roles. Esta secuencia es una propuesta de organización, no fechas de visita ya confirmadas. Al inicio se designan enlaces y se elige aplicación o Excel. Durante el mes se visita, recopila y corrige en paralelo. En la última semana se valida la cobertura de todos los servicios y la relación de colaboradores y roles. Se procura cerrar sin pendientes; lo que no pueda resolverse se registra con responsable y fecha. Analizar también la información académica disponible; el inicio del piloto se acordará con cada Dirección según las condiciones reales.')
for x,t,b in [(.65,'Inicio del mes','Acordar enlaces, agenda y herramienta de captura.'),(3.75,'Todo octubre','Visitar servicios y recopilar colaboradores y roles.'),(6.85,'En paralelo','Revisar datos, resolver dudas y analizar fuentes de alumnos.'),(9.95,'24–31 octubre','Validar información y procesos; documentar pendientes con fecha.')]:card(x,2.9,2.75,t,b,2.85)
banner('Meta al 31 de octubre: servicios visitados, datos validados y seguimiento de pendientes.')
new('Alumnos: aprovechar la información existente','Analizaremos las opciones de carga desde las aplicaciones académicas institucionales.','No pedir recapturar de entrada los datos de alumnos que ya existen en sistemas institucionales. Primero identificar aplicaciones, responsables y datos disponibles. Revisar opciones autorizadas de exportación a Excel o CSV, intercambio de archivos o integración técnica si el sistema lo permite. Definir los campos necesarios: matrícula, nombre, programa, semestre o grupo, ciclo y escuela. La matrícula permitirá cotejar registros y evitar duplicados; asignación a servicios y supervisores requiere validación aparte. Acordar permisos de acceso y actualización, realizar una prueba acotada y validar resultados antes de una carga general. Durante octubre se analizarán las opciones; no se promete que exista ya una conexión automática ni una fecha de integración sin conocer los sistemas.')
card(.65,2.75,3.9,'1 · Identificar','¿En qué aplicaciones están?\n¿Quién administra los datos?\n¿Qué información se puede obtener?',2.9)
card(4.72,2.75,3.9,'2 · Opciones','Exportación Excel o CSV.\nIntercambio de archivos.\nIntegración, si el sistema lo permite.',2.9)
card(8.79,2.75,3.9,'3 · Probar y validar','Cotejar matrícula y ciclo.\nEvitar duplicados.\nAcordar responsables de actualización.',2.9)
banner('Primero analizar y probar; después acordar la carga general y su calendario.')
new('Expedientes: contenido a revisar por cada Escuela','Se entregan dos documentos, con una explicación breve de cada campo.','Mostrar los PDF separados. Odontología tiene 56 campos o grupos comentados; Salud 90. Son propuestas de contenido y no una certificación de que el expediente clínico esté listo. Psicología se incluye como servicio de la Escuela de Salud, vinculado a La Casita y documentado en el directorio CPREVI 2024 de SEGEY. Confirmar con Dirección responsables, horarios y modalidades actuales, sin duplicar atenciones entre Psicología y Atención Comunitaria. Fuente: https://educacion.yucatan.gob.mx/prevencion/assets/archivos/DIRECTORIO_INSTANCIAS_CPREVI_2024.pdf. No discutir datos identificables de pacientes en esta reunión.')
card(.7,2.8,5.8,'Odontología','Historia, exploración y odontograma.\nPlan, consentimiento y evolución.\nPrácticas del alumno y supervisión.',2.65)
card(6.75,2.8,5.8,'Servicios de Salud','Psicología · La Casita incluida.\nCampos comunes y por disciplina.\nSeguimiento académico vinculado.',2.65)
banner('Cada servicio designa a quien revisará sus campos y documentos.')
new('Qué está preparado y qué falta acordar','La demostración permite revisar el enfoque antes de operar con información real.','Evitar presentar el alcance clínico completo como terminado. Están preparados la demostración, paneles académicos, plantilla Excel y propuestas documentales. Faltan validación institucional y los desarrollos clínicos y de importación correspondientes. Antes de datos reales deben revisarse accesos y reactivarse las medidas de seguridad suspendidas para la demostración.')
card(.65,2.75,3.9,'Preparado','Aplicación de demostración.\nPaneles académicos.\nExcel y propuestas de expedientes.',2.85)
card(4.72,2.75,3.9,'Por validar','Procesos y responsables.\nCampos, permisos y alcance clínico.\nCarga de Excel y fuentes de alumnos.',2.85)
card(8.79,2.75,3.9,'Acompañamiento','Visitas durante todo octubre.\nEscucha y revisión a conciencia.\nCierre y seguimiento de pendientes.',2.85)
banner('La evaluación y las decisiones clínicas permanecen bajo responsabilidad profesional.')
new('Cerremos con responsables y fechas','Seis acuerdos para trabajar durante octubre y revisar el cierre el día 31.','Completar los acuerdos: enlace por Escuela y herramienta elegida por cada Dirección; visitas y al menos dos días por servicio distribuidos durante octubre; responsable de identificar sistemas y analizar opciones de carga de alumnos; servicio y condiciones del piloto; revisores de expedientes; reunión de seguimiento y revisión al 31 de octubre. La relación de colaboradores y roles comparte el plazo de octubre. Procurar completar el trabajo a conciencia y sin pendientes; documentar las excepciones con responsable y fecha. Registrar como pendientes los acuerdos aún no confirmados.')
items=[('1','Colaboradores y roles','Aplicación o Excel; revisión y cierre al 31 de octubre.'),('2','Visitas de trabajo','Todo octubre; mínimo dos días por servicio y seguimiento.'),('3','Información de alumnos','Responsable y sistemas a revisar; analizar opciones de carga.'),('4','Piloto','Servicio, equipo y condiciones para iniciar.'),('5','Expedientes','Revisor de cada servicio y fecha de devolución.'),('6','Cierre y seguimiento','Revisión al 31 de octubre; pendientes con responsable y fecha.')]
for i,(n,t,b) in enumerate(items):
 y=2.5+i*.68;rect(.7,y,.42,.42,TEAL);txt(.82,y+.055,.3,.3,n,15,WHITE,True);txt(1.32,y,3.9,.45,t,17,INK,True);txt(5.25,y+.02,7.1,.5,b,16)
slides.append(''.join(parts))
r.save(O/'Presentacion_Directores_Salud_Odontologia.pptx')
html='<!doctype html><html lang="es"><meta charset="utf-8"><title>Reunión con Direcciones</title><style>@page{size:13.333in 7.5in;margin:0}*{box-sizing:border-box}body{margin:0;font-family:Arial,sans-serif}section{position:relative;width:13.333in;height:7.5in;background:#'+BG+';overflow:hidden;break-after:page}section:last-child{break-after:auto}</style><body>'+''.join('<section>'+p+'</section>' for p in slides)+'</body></html>'
(O/'Presentacion_Directores_Salud_Odontologia.html').write_text(html)
(O/'Guion_del_presentador.txt').write_text('GUION DE APOYO · REUNIÓN CON DIRECCIONES\nDuración sugerida: 35–40 minutos, incluyendo acuerdos.\n\n'+'\n'.join(notes))
print('Diapositivas:',len(r.slides))
