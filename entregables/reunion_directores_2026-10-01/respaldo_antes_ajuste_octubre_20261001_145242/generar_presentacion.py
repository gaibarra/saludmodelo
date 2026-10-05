from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
from html import escape
import base64
from PIL import Image
O=Path('/home/gaibarra/plandetrabajo/entregables/reunion_directores_2026-10-01');SRC=O.parent/'reunion_directores_2026-09-30'
r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
BG='F3F9F8';INK='173F49';TEAL='087F82';MUT='54717B';WHITE='FFFFFF';PALE='E2F1ED';GOLD='FFF1CB'
slides=[];notes=[];s=None;parts=[]
def rect(x,y,w,h,color):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
 parts.append(f'<div style="position:absolute;left:{x}in;top:{y}in;width:{w}in;height:{h}in;background:#{color}"></div>')
def txt(x,y,w,h,text,size=22,color=INK,bold=False):
 sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(8)
 parts.append(f'<div style="position:absolute;left:{x}in;top:{y}in;width:{w}in;height:{h}in;font-size:{size}pt;color:#{color};font-weight:{700 if bold else 400};line-height:1.18;white-space:pre-line">{escape(text)}</div>')
def new(title,kicker,note):
 global s,parts
 if s is not None:slides.append(''.join(parts))
 s=r.slides.add_slide(r.slide_layouts[6]);parts=[];s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 rect(0,0,.16,7.5,TEAL);txt(.55,.32,12,.3,'UNIVERSIDAD MODELO  /  SALUD Y ODONTOLOGÍA',10,TEAL,True)
 txt(.55,.87,12.1,1,title,30,INK,True);txt(.58,1.9,12,.5,kicker,16,MUT)
 txt(.58,7.12,11,.2,'REUNIÓN CON DIRECCIONES · PROPUESTA DE TRABAJO · 01 OCT 2026',9,MUT)
 txt(12.15,7.1,.6,.25,f'{len(r.slides):02d}',11,TEAL,True)
 s.notes_slide.notes_text_frame.text=note;notes.append(f'{len(r.slides):02d}. {title}\n{note}\n')
def card(x,y,w,title,body,h=2.3):
 rect(x,y,w,h,WHITE);rect(x,y,.055,h,TEAL);txt(x+.2,y+.2,w-.4,.6,title,21,TEAL,True);txt(x+.2,y+(.68 if h<1.8 else .98),w-.4,h-(.74 if h<1.8 else 1.05),body,17 if h<1.8 else 18)
def banner(text,y=6.3):rect(.58,y,12.1,.57,PALE);txt(.78,y+.12,11.7,.4,text,14,TEAL,True)
def photo(name,x,y,w,h):
 path=(O/'imagenes'/name) if (O/'imagenes'/name).exists() else SRC/name;im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale
 x+=(w-nw)/2;y+=(h-nh)/2;s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(nw),height=Inches(nh))
 b=base64.b64encode(path.read_bytes()).decode();parts.append(f'<img src="data:image/png;base64,{b}" style="position:absolute;left:{x}in;top:{y}in;width:{nw}in;height:{nh}in">')
new('¿Qué es la aplicación Salud y Odontología?','Una plataforma web para organizar servicios y acompañar la formación de los alumnos.','Presentar la aplicación como una plataforma institucional que conecta la atención de los servicios, la organización del personal y la formación académica. Participan pacientes y usuarios, colaboradores, alumnos, supervisores y Direcciones, cada uno con permisos específicos. El alcance previsto incluye solicitudes de atención, seguimiento y expedientes por servicio, vinculación de prácticas con competencias e informes individuales y consolidados. Las escuelas conservan su administración independiente. Aclarar que el despliegue es por etapas: la demostración y el núcleo académico sirven de base, mientras que la operación clínica real requiere validación y completar los módulos correspondientes.')
txt(.7,2.55,11.9,1.22,'Conecta a pacientes y usuarios, colaboradores, alumnos, supervisores y Direcciones para dar seguimiento al servicio prestado y al aprendizaje que genera.',23,INK)
card(.65,3.85,3.9,'Atención y servicios','Solicitudes de atención, seguimiento y expedientes adaptados a cada servicio.',2.15)
card(4.72,3.85,3.9,'Formación','Prácticas, supervisión, competencias e informes por alumno y ciclo.',2.15)
card(8.79,3.85,3.9,'Gestión y Dirección','Responsables, tareas y avances; cada Escuela con administración propia.',2.15)
banner('Implementación por etapas: demostración y núcleo académico; alcance clínico por validar y completar.')
new('El objetivo: conocer y respaldar cada práctica','Una visión compartida para las Direcciones de Salud y Odontología.','Abrir con el objetivo académico. Preguntar qué información resulta indispensable para cada Dirección durante el ciclo y al finalizarlo. No presentar las cifras simuladas como resultados reales.')
txt(.7,2.7,11.7,1.7,'Saber qué hizo cada alumno,\nquién lo supervisó y qué competencias alcanzó.',32,INK,True)
card(.7,4.7,5.8,'Durante el ciclo','Avances, pendientes y seguimiento oportuno.',1.35)
card(6.7,4.7,5.8,'Al finalizar el ciclo','Informes individuales y consolidados.',1.35)
banner('Hoy buscamos acuerdos para comenzar con personas y procesos reales.')
new('Lo que Dirección necesita poder responder','Primero acordamos la información útil; después afinamos los formularios.','Pedir a los Directores que prioricen estas cuatro preguntas y agreguen lo que falte. Acordar qué información será necesaria para dar por cerrado el ciclo académico.')
card(.65,2.7,5.9,'01 · Participación','¿Qué prácticas realizó cada alumno y en qué servicio?',1.5)
card(6.75,2.7,5.9,'02 · Supervisión','¿Quién revisó y validó cada actividad?',1.5)
card(.65,4.45,5.9,'03 · Aprendizaje','¿Qué competencias demostró y cuáles debe reforzar?',1.5)
card(6.75,4.45,5.9,'04 · Cierre','¿Qué resultados y pendientes deja el ciclo?',1.5)
new('Una aplicación que conecta el trabajo cotidiano','Del registro de una actividad al informe académico.','Explicar el flujo en lenguaje sencillo. Una captura no equivale automáticamente a una práctica acreditada: debe existir revisión del supervisor. Los módulos académicos sirven de base; los formatos clínicos siguen sujetos a validación y desarrollo.')
for x,title,body in [( .65,'1 · Registrar','Actividad, alumno, servicio y fecha.'),(3.75,'2 · Supervisar','Revisión, observaciones y correcciones.'),(6.85,'3 · Validar','Actividad y competencia acreditadas.'),(9.95,'4 · Consultar','Avances e informes para Dirección.')]:card(x,2.9,2.75,title,body,2.6)
banner('La validación del supervisor convierte el registro en evidencia académica.')
new('Dos escuelas, administración independiente','Salud y Odontología tienen el mismo nivel jerárquico.','Confirmar la separación administrativa. Cada Dirección organiza su escuela, sus responsables y el seguimiento de sus alumnos.')
card(.7,2.8,5.8,'Escuela de Salud','Su Dirección administra sus servicios, colaboradores y seguimiento académico.',2.3)
card(6.75,2.8,5.8,'Escuela de Odontología','Su Dirección administra su equipo, operación y prácticas odontológicas.',2.3)
txt(.9,5.5,11.5,.7,'Cada Escuela define sus responsables y administra su información,\ncon funciones y permisos propios.',20)
new('Así se podrá consultar el avance · Salud','Vista ilustrativa de la interfaz actual con cifras ficticias.','Todas las cifras de esta imagen son simuladas y corresponden únicamente al ejemplo de la Escuela de Salud. No se cargaron registros ficticios en producción para estas imágenes.')
photo('02_Direccion_Salud.png',.6,2.48,8.5,4.38)
txt(9.3,2.9,3.3,3.4,'Administración de Salud.\n\nPrácticas supervisadas.\n\nAvance de sus alumnos.',21,INK,True)
new('Así se podrá consultar el avance · Odontología','Vista ilustrativa de la interfaz actual con cifras ficticias.','Se muestra la Escuela de Odontología con administración propia. Las cifras ilustran el funcionamiento posterior, no los registros reales ni una meta institucional. Los paneles completos también se entregan como PNG y PDF aparte.')
photo('03_Direccion_Odontologia.png',.6,2.48,8.5,4.38)
txt(9.3,2.9,3.3,3.4,'Administración propia.\n\nPrácticas supervisadas.\n\nInformes del ciclo.',21,INK,True)
new('Primer compromiso: colaboradores en siete días','Cada Escuela entrega una relación revisada por su Dirección.','Mostrar el Excel preparado. Una fila por colaborador y servicio; repetir correo si participa en varios servicios. La plantilla está lista; la importación no es automática y debe prepararse y validarse. No pedir contraseñas ni información clínica en ese archivo.')
card(.65,2.75,3.9,'Qué registrar','Nombre y correo institucional.\nServicio y función.\nParticipación y vigencia.',2.8)
card(4.72,2.75,3.9,'Cómo colaborar','Captura en la aplicación o plantilla Excel.\nUn responsable por Escuela.',2.8)
card(8.79,2.75,3.9,'Cómo entregar','Dirección revisa la relación.\nSe consolidan las copias y se revisan duplicados.',2.8)
banner('Excel es una alternativa de recopilación; su carga requiere revisión e importación preparada.')
new('Mi compromiso: al menos dos días por servicio','C.P. Gonzalo Arturo Ibarra Mendoza · Presencia física y escucha del equipo.','Expresar el compromiso personal de acompañamiento. Las visitas no son una auditoría del personal. Buscan comprender el proceso actual, sus excepciones y los formatos utilizados. Coordinar horarios sin interferir con la atención y respetar la confidencialidad. Dos días mínimos por servicio; ampliar si hace falta.')
card(.7,2.8,5.8,'Día 1 · Conocer y escuchar','Recorrer el proceso actual.\nEscuchar a colaboradores y supervisores.\nIdentificar formatos y dificultades.',2.65)
card(6.75,2.8,5.8,'Día 2 · Revisar y acordar','Revisar casos y excepciones.\nValidar lo observado con el equipo.\nAcordar ajustes y prioridades.',2.65)
banner('Un enlace por servicio y calendario acordado. Las visitas pueden ampliarse.')
new('Calendario propuesto: avanzar en paralelo','Los siete días corresponden a la recopilación; las visitas tienen su propio calendario.','No prometer recorrer todos los servicios en una semana. Designar un enlace y reservar al menos dos jornadas por servicio. En la reunión se completa la fecha exacta de entrega del registro: siete días después de la reunión. La fecha del piloto depende de responsables, alcance y condiciones de operación.')
for x,t,b in [(.65,'En la reunión','Acordar responsables, servicio piloto y fechas.'),(3.75,'Días 1–4','Recopilar colaboradores en aplicación o Excel.'),(6.85,'Días 5–7','Revisar, corregir y entregar la relación.'),(9.95,'Visitas','Dos días por servicio, según agenda acordada.')]:card(x,2.9,2.75,t,b,2.6)
banner('Después: validar los ajustes y acordar el inicio del piloto; Odontología es la propuesta inicial.')
new('Expedientes: contenido a revisar por cada Escuela','Se entregan dos documentos, con una explicación breve de cada campo.','Mostrar los PDF separados. Odontología tiene 56 campos o grupos comentados; Salud 90. Son propuestas de contenido y no una certificación de que el expediente clínico esté listo. Psicología se incluye como servicio de la Escuela de Salud, vinculado a La Casita y documentado en el directorio CPREVI 2024 de SEGEY. Confirmar con Dirección responsables, horarios y modalidades actuales, sin duplicar atenciones entre Psicología y Atención Comunitaria. Fuente: https://educacion.yucatan.gob.mx/prevencion/assets/archivos/DIRECTORIO_INSTANCIAS_CPREVI_2024.pdf. No discutir datos identificables de pacientes en esta reunión.')
card(.7,2.8,5.8,'Odontología','Historia, exploración y odontograma.\nPlan, consentimiento y evolución.\nPrácticas del alumno y supervisión.',2.65)
card(6.75,2.8,5.8,'Servicios de Salud','Psicología · La Casita incluida.\nCampos comunes y por disciplina.\nSeguimiento académico vinculado.',2.65)
banner('Cada servicio designa a quien revisará sus campos y documentos.')
new('Qué está preparado y qué falta acordar','La demostración permite revisar el enfoque antes de operar con información real.','Evitar presentar el alcance clínico completo como terminado. Están preparados la demostración, paneles académicos, plantilla Excel y propuestas documentales. Faltan validación institucional y los desarrollos clínicos y de importación correspondientes. Antes de datos reales deben revisarse accesos y reactivarse las medidas de seguridad suspendidas para la demostración.')
card(.65,2.75,3.9,'Preparado','Aplicación de demostración.\nPaneles académicos.\nExcel y propuestas de expedientes.',2.85)
card(4.72,2.75,3.9,'Por validar','Procesos y responsables.\nCampos y permisos.\nImportación de Excel y alcance clínico.',2.85)
card(8.79,2.75,3.9,'Acompañamiento','Visitas a cada servicio.\nEscucha de los equipos.\nAjustes a los procesos reales.',2.85)
banner('La evaluación y las decisiones clínicas permanecen bajo responsabilidad profesional.')
new('Cerremos con responsables y fechas','Cinco acuerdos concretos para pasar de la demostración al trabajo institucional.','Completar estos acuerdos durante la reunión: responsable de recopilación de cada Escuela; fecha de entrega en siete días; enlace y calendario de visitas por servicio; servicio piloto y sus participantes; revisores de expedientes y fecha de seguimiento. Registrar acuerdos pendientes como pendientes, no como autorizaciones obtenidas.')
items=[('1','Registro de colaboradores','Responsable por Escuela y entrega en siete días.'),('2','Visitas de trabajo','Enlace y al menos dos días presenciales por servicio.'),('3','Piloto','Servicio, equipo y condiciones para iniciar.'),('4','Expedientes','Revisor de cada servicio y fecha de devolución.'),('5','Seguimiento','Fecha de próxima reunión y acuerdos por comprobar.')]
for i,(n,t,b) in enumerate(items):
 y=2.55+i*.73;rect(.7,y,.45,.45,TEAL);txt(.83,y+.065,.3,.3,n,16,WHITE,True);txt(1.35,y,4,.4,t,18,INK,True);txt(5.3,y+.03,7,.4,b,18)
slides.append(''.join(parts))
r.save(O/'Presentacion_Directores_Salud_Odontologia.pptx')
html='<!doctype html><html lang="es"><meta charset="utf-8"><title>Reunión con Direcciones</title><style>@page{size:13.333in 7.5in;margin:0}*{box-sizing:border-box}body{margin:0;font-family:Arial,sans-serif}section{position:relative;width:13.333in;height:7.5in;background:#'+BG+';overflow:hidden;break-after:page}section:last-child{break-after:auto}</style><body>'+''.join('<section>'+p+'</section>' for p in slides)+'</body></html>'
(O/'Presentacion_Directores_Salud_Odontologia.html').write_text(html)
(O/'Guion_del_presentador.txt').write_text('GUION DE APOYO · REUNIÓN CON DIRECCIONES\nDuración sugerida: 35–40 minutos, incluyendo acuerdos.\n\n'+'\n'.join(notes))
print('Diapositivas:',len(r.slides))
