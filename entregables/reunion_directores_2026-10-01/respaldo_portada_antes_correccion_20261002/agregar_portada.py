"""Añade una portada editable, conservando las 14 diapositivas originales.
Ejecutar con PYTHONPATH=.cache/reunion-slides/site-packages python3.
El respaldo es la entrada estable: repetir no duplica la portada.
"""
from pathlib import Path
import shutil
import base64
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from html import escape
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from pypdf import PdfReader, PdfWriter

O=Path(__file__).resolve().parents[1]
B=O/'respaldo_antes_portada'
B.mkdir(exist_ok=True)
for name in ['Presentacion_Directores_Salud_Odontologia.pptx','Presentacion_Directores_Salud_Odontologia.pdf','Presentacion_Directores_Salud_Odontologia.html','Guion_del_presentador.txt']:
    if not (B/name).exists(): shutil.copy2(O/name,B/name)
logo=O/'fuentes/portada/Modelo.jpg'
if not logo.exists(): shutil.copy2(O.parents[1]/'frontend/public/images/Modelo.jpg',logo)
pdf=canvas.Canvas(str(O/'fuentes/portada/portada.pdf'),pagesize=(13.333*72,7.5*72))
r=Presentation(B/'Presentacion_Directores_Salud_Odontologia.pptx')
s=r.slides.add_slide(r.slide_layouts[6]);parts=[]
s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('FFFFFF')
def rect(x,y,w,h,color):
    pdf.setFillColor(HexColor("#"+color));pdf.rect(x*72,(7.5-y-h)*72,w*72,h*72,stroke=0,fill=1)
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
    parts.append(f'<div style="position:absolute;left:{x}in;top:{y}in;width:{w}in;height:{h}in;background:#{color}"></div>')
def txt(x,y,w,h,t,size,color='163B51',bold=False):
    pdf.setFillColor(HexColor('#'+color));pdf.setFont('Helvetica-Bold' if bold else 'Helvetica',size)
    for i,line in enumerate(t.split('\n')):pdf.drawString(x*72,(7.5-y)*72-size*.93-i*size*1.18,line)
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0)
    parts.append(f'<div style="position:absolute;left:{x}in;top:{y}in;width:{w}in;height:{h}in;font-family:Arial;font-size:{size}pt;line-height:1.18;color:#{color};font-weight:{700 if bold else 400};white-space:pre-line">{escape(t)}</div>')
rect(0,0,.17,7.5,'087F82')
rect(0,7.29,13.333,.21,'163B51')
s.shapes.add_picture(str(logo), Inches(.64), Inches(.44), width=Inches(1.10))
pdf.drawImage(str(logo),.64*72,(7.5-.44-1.10*225/224)*72,width=1.10*72,height=1.10*225/224*72)
b=base64.b64encode(logo.read_bytes()).decode()
parts.append(f'<img src="data:image/jpeg;base64,{b}" style="position:absolute;left:.64in;top:.44in;width:1.10in">')
txt(2.0,.60,10.5,.45,'UNIVERSIDAD MODELO',25,bold=True)
txt(2.02,1.12,10,.35,'ESCUELA DE SALUD  /  ESCUELA DE ODONTOLOGÍA',13,'54717B')
rect(.7,1.85,11.93,.02,'D7E5E7')
txt(.7,2.13,11,.3,'REUNIÓN CON DIRECCIONES',13,'087F82',True)
txt(.7,2.66,12,1.35,'Salud y Odontología\nPlan de trabajo institucional',35,bold=True)
txt(.73,4.04,11.6,.45,'Organización de servicios y seguimiento de la formación académica',18,'54717B')
rect(.7,4.91,5.82,1.20,'F0F6F6');rect(6.75,4.91,5.88,1.20,'F0F6F6')
txt(.92,5.12,5.4,.38,'Carla Amira Leon Pinto',23,bold=True)
txt(.94,5.67,5.3,.30,'Directora de la Escuela de Salud',13,'54717B')
txt(6.97,5.12,5.5,.38,'C.D. E.P. Mario Sosa Correa',22,bold=True)
txt(6.99,5.67,5.4,.30,'Director de la Escuela de Odontología',13,'54717B')
txt(.72,6.55,8,.27,'PRESENTA · C.P. Gonzalo Arturo Ibarra Mendoza',13,bold=True)
txt(.72,6.94,10,.25,'Propuesta de trabajo · Octubre de 2026',11,'54717B')
txt(9.32,6.55,3.5,.35,'1 de octubre de 2026',16,'087F82',True)
s.notes_slide.notes_text_frame.text='Portada institucional. Nombres tomados de docs/DECISIONES.md, decisiones 125 y 165. Escudo original conservado. Fecha de la reunión: 1 de octubre de 2026.'
ids=r.slides._sldIdLst;node=ids[-1];ids.remove(node);ids.insert(0,node)
r.save(O/'Presentacion_Directores_Salud_Odontologia.pptx')
cover='<section style="background:white">'+''.join(parts)+'</section>'
original=(B/'Presentacion_Directores_Salud_Odontologia.html').read_text()
(O/'Presentacion_Directores_Salud_Odontologia.html').write_text(original.replace('<body>','<body>'+cover,1))
style='<style>@page{size:13.333in 7.5in;margin:0}*{box-sizing:border-box}body{margin:0}section{position:relative;width:13.333in;height:7.5in;overflow:hidden}</style>'
(O/'fuentes/portada/portada.html').write_text('<!doctype html><html lang="es"><meta charset="utf-8">'+style+'<body>'+cover+'</body></html>')
print('PPTX y HTML actualizados; originales conservados en',B)

pdf.showPage();pdf.save()
w=PdfWriter();w.append(O/'fuentes/portada/portada.pdf');w.append(B/'Presentacion_Directores_Salud_Odontologia.pdf')
with open(O/'Presentacion_Directores_Salud_Odontologia.pdf','wb') as f:w.write(f)
