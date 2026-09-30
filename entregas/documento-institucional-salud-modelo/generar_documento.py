from pathlib import Path
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT

OUT=Path(__file__).resolve().parent
pages=[]
def page(title, subtitle=''):
 p={'title':title,'subtitle':subtitle,'items':[]}; pages.append(p); return p

def para(p,t):p['items'].append(('p',t))
def head(p,t):p['items'].append(('h',t))
def table(p,headers,rows):p['items'].append(('t',(headers,rows)))

p=page('Salud Modelo','Exposición de motivos y plan de participación y adopción')
para(p,'Documento para la Dirección de la Escuela de Salud y los colaboradores de sus servicios. Propuesta para revisión y acuerdo institucional • 28 de septiembre de 2026.')
head(p,'1. ¿Por qué crear esta aplicación?')
para(p,'La Escuela de Salud reúne servicios con necesidades propias y actividades que se relacionan entre sí: recepción, atención profesional, seguimiento, enseñanza, administración y uso de recursos. Para construir una aplicación que les resulte útil, primero es necesario comprender cómo trabajan, qué documentos utilizan, quién toma cada decisión y qué sucede cuando un caso se aparta de lo habitual.')
para(p,'El plan de trabajo plantea documentar y validar esos procesos con quienes los realizan. El gestor de planificación Salud Modelo se crea para organizar ese esfuerzo en un espacio común: reunir respuestas y documentos, aclarar dudas, asignar tareas y dar seguimiento a los acuerdos. Su objetivo central es convertir el conocimiento de los servicios en un plan de trabajo compartido, verificable y realizable.')
para(p,'Sin un mecanismo común, existe el riesgo de que las respuestas queden dispersas en conversaciones o archivos, se interpreten de manera diferente o se conviertan en requisitos sin confirmación. La aplicación busca prevenir esos problemas. No se afirma que todos ellos ocurran hoy en la Escuela; el levantamiento permitirá identificar cuáles existen y qué mejoras tienen prioridad.')
para(p,'La participación de los colaboradores es indispensable: ellos conocen las actividades cotidianas, las dificultades y las soluciones que ya funcionan. La tecnología organiza esa información; los responsables del servicio confirman su validez y Dirección decide prioridades, recursos y cambios.')
head(p,'Resultados que se buscan')
table(p,['Para quién','Resultado esperado'],[
['Dirección','Contar con una visión comprensible de avances, pendientes, decisiones y necesidades de apoyo.'],
['Responsables de servicio','Tener procesos acordados, tareas con responsables y fechas compatibles con la disponibilidad del equipo.'],
['Colaboradores','Saber qué deben aportar, dónde hacerlo, quién revisará su trabajo y cómo resolver una duda.'],
['Proyecto institucional','Construir los siguientes componentes de la aplicación sobre información confirmada, conservando la explicación de cada decisión.']])
para(p,'La primera decisión se abordará en la reunión por Teams del 29 de septiembre de 2026, a las 10:00, con la directora Carla Amira Leon Pinto: definir el servicio piloto, sus participantes y las condiciones para probar el gestor.')

p=page('Qué abarca esta etapa','El gestor como primer paso del proyecto institucional')
head(p,'2. Qué permitirá hacer el gestor')
para(p,'El gestor permite organizar preguntas por servicio, consultar orientaciones, guardar respuestas y adjuntar documentos autorizados. Facilita que otra persona revise lo aportado, solicite aclaraciones o valide el contenido. También organiza tareas, fechas, responsables, suplencias, disponibilidad y ausencias, además de decisiones, informes y avisos dentro de la aplicación.')
para(p,'Al planificar una tarea, el equipo debe indicar el resultado que espera, cuánto trabajo estima y quién puede realizarlo. Si una ausencia o una carga excesiva impide cumplir, se revisa el compromiso y se tramita el ajuste correspondiente. Una suplencia permite continuidad con permisos y vigencia definidos; los cambios de responsabilidad o calendario deben quedar explícitos.')
head(p,'Relación con la aplicación completa')
para(p,'El plan original tiene un alcance mayor: documentar procesos y desarrollar una aplicación que conecte admisión, atención, seguimiento, docencia y administración, además de una página institucional para conocer servicios y solicitar citas. Ese alcance incluye expedientes por especialidad, reglas de acceso y otros componentes que requieren desarrollo y validación propios.')
para(p,'El gestor es la herramienta para coordinar esa construcción. Su desarrollo y sus pruebas técnicas no significan que los módulos clínicos estén terminados ni que exista autorización para operar con pacientes. A la fecha de este documento, el gestor requiere piloto humano, acuerdos de operación y aceptación institucional; no se presenta como un sistema ya desplegado en producción.')
head(p,'Palabras que usaremos')
table(p,['Término','Significado sencillo'],[
['Proceso','Secuencia de actividades desde que se recibe una solicitud hasta que se obtiene un resultado.'],
['Evidencia','Formato, documento o registro autorizado que permite comprobar o explicar una respuesta.'],
['Validar','Revisar que lo registrado corresponde a la operación y aprobar expresamente esa versión.'],
['Piloto','Ensayo con un grupo y un alcance pequeños para aprender y corregir antes de ampliar el uso.'],
['Excepción','Situación distinta del caso habitual, como la ausencia de un responsable o un documento faltante.'],
['Dependencia','Trabajo que debe resolverse antes de que otra tarea pueda avanzar.']])
head(p,'Uso responsable de la información y de la IA')
para(p,'En el piloto se usarán datos ficticios o ejemplos anonimizados autorizados. Cada persona accederá a la información que corresponda a su función. Si se habilita asistencia de inteligencia artificial, será bajo una política aprobada: podrá apoyar la redacción o comprensión, pero sus sugerencias deberán revisarse. No decidirá diagnósticos, autorizaciones profesionales ni cumplimiento institucional.')

p=page('Cómo participa una persona por primera vez','Guía de trabajo sin conocimientos técnicos')
head(p,'3. Recorrido del usuario')
table(p,['Paso','Qué hace la persona','Qué obtiene'],[
['1. Recibir orientación','Asiste a la explicación inicial y confirma su servicio, función y persona de apoyo. Recibe el acceso por el procedimiento autorizado.','Sabe qué le corresponde y a quién consultar.'],
['2. Revisar su encargo','Consulta las preguntas o tareas asignadas y lee la orientación disponible.','Entiende qué información debe aportar.'],
['3. Describir el trabajo real','Explica cómo se realiza hoy la actividad: quién inicia, qué se hace, qué documentos se usan y cómo termina.','Una descripción comprensible de la operación actual.'],
['4. Aportar respaldo','Adjunta formatos vacíos o ejemplos autorizados e indica su procedencia. Guarda el avance para continuar después.','Información que otra persona puede revisar.'],
['5. Registrar dudas','Indica lo que desconoce o falta confirmar; avisa al responsable para convertirlo en un pendiente con seguimiento.','La duda tiene un camino de resolución.'],
['6. Enviar a revisión','Solicita revisión cuando lo aportado está listo. Atiende las observaciones y conserva las correcciones.','Una respuesta validada o una aclaración concreta por resolver.'],
['7. Acordar tareas','Participa en la definición de resultados, fechas y esfuerzo; comunica su disponibilidad y ausencias por el procedimiento acordado.','Compromisos que pueden revisarse y cumplirse.'],
['8. Dar seguimiento','Actualiza avances, registra el trabajo realizado cuando corresponda y consulta los avisos. Confirma la lectura de los informes recibidos.','El equipo conoce el estado real del trabajo.']])
head(p,'Qué hacer cuando algo no está claro')
para(p,'Si no conoce una respuesta, no la invente. Explique qué falta y quién podría aclararlo. Si un punto parece no aplicar, indique el motivo para que el responsable lo revise. Si dos personas describen reglas diferentes, registre la diferencia: el equipo debe acordar cuál es la regla vigente antes de programarla.')
para(p,'Una devolución significa que hace falta aclarar o completar algo. Lea el motivo, corrija lo necesario y vuelva a enviar. Si no puede entrar o no encuentra su asignación, comuníquelo al enlace designado. En la reunión inicial se confirmarán el nombre de ese enlace y el medio de atención; todavía no están definidos.')

p=page('Ejemplo y organización cotidiana','Cómo convertir una duda en un resultado acordado')
head(p,'4. Un ejemplo sencillo')
para(p,'Ejemplo ficticio para explicar el método: un servicio necesita documentar qué ocurre cuando la persona encargada de confirmar citas está ausente. Este ejemplo no describe una regla ya aprobada de la Escuela.')
para(p,'La colaboradora asignada explica cómo se confirma una cita y aporta un formato vacío. Al responder sobre ausencias, detecta que falta confirmar quién puede sustituir al titular. Registra la duda y la comunica al responsable; no da por hecho que cualquier compañero puede asumir esa función.')
para(p,'El responsable consulta al equipo y propone una regla de suplencia. Una persona competente revisa los permisos y el procedimiento. Una vez acordados, se crea una tarea para documentar la regla y comprobarla, con responsable, fecha y resultado esperado: poder explicar quién actúa, durante qué periodo y qué registro deja.')
para(p,'Si la persona asignada no tiene disponibilidad suficiente, se propone ajustar fechas o distribución del trabajo y se obtiene la aprobación correspondiente. Durante el piloto se prueba el caso normal y las excepciones acordadas. El resultado se registra; si falla, se corrige y se repite la prueba afectada antes de aceptar el proceso.')
head(p,'Ritmo de colaboración propuesto')
table(p,['Momento','Actividad','Participan'],[
['Inicio del trabajo','Inducción práctica de 45–60 minutos con un ejemplo y oportunidad de intentar el recorrido.','Participantes del piloto y facilitador designado.'],
['Durante la semana','Bloques de trabajo acordados para responder, reunir documentos y resolver observaciones.','Colaboradores según su asignación y disponibilidad.'],
['Una vez por semana','Revisión de 30 minutos: avances, bloqueos, próximas tareas y necesidad de apoyo.','Responsable, suplente y colaboradores involucrados.'],
['Al llegar a un hito','Revisar productos y decidir si se aceptan, corrigen o reprograman.','Responsable del servicio, revisor y Dirección cuando corresponda.']])
para(p,'Estas duraciones son una propuesta de organización, no una carga autorizada. Cada responsable deberá reservar tiempo de colaboración compatible con la atención y las demás funciones del servicio. La disponibilidad se confirma antes de comprometer fechas.')
head(p,'Cómo evitar trabajo innecesario')
para(p,'Se comienza por los procesos críticos del piloto y se aprovechan los formatos vigentes. No es necesario contestar todo el inventario el primer día. La captura debe concentrarse en información útil para tomar decisiones. El plan original contempla entrevistas de 90 minutos, dos observaciones por unidad y una validación de 60 minutos; su programación se acordará con cada servicio.')

p=page('Quién participa y qué aporta','Responsabilidades compartidas, con decisiones identificables')
head(p,'5. Organización de la participación')
para(p,'Dirección confirma el alcance, designa a quien coordina el proyecto, resuelve prioridades y autoriza los hitos que le correspondan. Carla Amira Leon Pinto participará en la definición del piloto como directora; la responsabilidad de ejecutarlo y los demás nombramientos se acordarán en la reunión.')
para(p,'Cada servicio designará un responsable y un suplente. El responsable organizará las aportaciones, resolverá dudas y coordinará la validación. Los colaboradores describirán las actividades que conocen y probarán los recorridos asignados. La revisión corresponderá a una persona competente distinta del autor cuando se requiera aprobación independiente. Gonzalo coordina análisis y desarrollo según la propuesta original, sujeto a la confirmación organizativa del arranque.')
table(p,['Servicio o área','Aportación principal de sus colaboradores'],[
['Odontología','Explicar valoración, odontograma, plan, presupuesto, supervisión, instrumental, esterilización e imágenes; aportar sus formatos y excepciones.'],
['Fisioterapia','Describir evaluación, escalas, objetivos, sesiones, paquetes y disponibilidad de terapeutas y equipos.'],
['Nutrición','Precisar reglas de ambas sedes, fórmulas, unidades, planes, seguimiento y tarifas vigentes.'],
['Psicología','Aclarar participantes, permisos, representación, consentimientos y procedimientos ante situaciones especiales, con revisión competente.'],
['Deporte y readaptación','Confirmar pruebas que realmente realizan, variables, reportes, equipos y derivaciones.'],
['UAPS y prevención','Precisar actividades autorizadas, registros, referencias y notificaciones que correspondan a su operación.'],
['Comunidad laboral y educativa / La Casita','Aportar convenios, jornadas, sedes, participantes, cuotas y reglas de uso autorizado de datos.'],
['Jurídico — inclusión por confirmar','Si se incorpora, describir representación, conflictos de interés, plazos y separación de expedientes.'],
['Administración, docencia y recursos','Acordar recepción, tarifas, saldos, reglas fiscales, rotaciones, supervisión, inventario y mantenimiento.']])
para(p,'Todas las unidades preparan información en paralelo desde el 1 de octubre. Cuando una actividad pasa de un servicio a otro, ambos revisan qué se entrega, quién lo recibe y cómo se confirma la recepción. No se atribuyen a una sola unidad reglas que afectan a otras.')

p=page('Calendario para lograr la adopción','Fechas de 2026 • propuesta de ejecución y productos verificables')
head(p,'6. Cronograma de trabajo del gestor')
para(p,'La reunión del 29 de septiembre es coordinación previa. El calendario de ejecución del plan original comienza el 1 de octubre. La siguiente propuesta organiza la adopción del gestor alrededor de sus hitos; requiere confirmar responsables y capacidad. Las fechas no constituyen aceptación ni garantía de entrega del alcance completo.')
table(p,['Fecha','Actividad y quién la conduce','Producto para verificar avance'],[
['29 sep., 10:00\nTeams','Dirección y coordinación: elegir piloto y acordar condiciones.','Servicio, responsable, suplente, participantes y alcance inicial definidos.'],
['1–2 oct.','Dirección y responsables: confirmar alcance, nombramientos y disponibilidad.','Lista de participantes, permisos y agenda de trabajo acordados.'],
['5 oct.','Coordinación y equipo piloto: presentar el gestor y realizar inducción.','Accesos comprobados y primer ejercicio guiado. Hito mínimo del calendario original.'],
['7 oct.','Todas las unidades: entregar formatos, catálogo, volúmenes y evidencia inicial.','Inventario inicial y faltantes identificados por unidad.'],
['9 oct.','Responsables y coordinación: validar procesos críticos y revisar esfuerzo.','Prioridades confirmadas y reestimación del alcance y calendario.'],
['16 oct.','Revisores de servicio: resolver excepciones y revisar orientaciones y reglas aplicables.','Primer conjunto de respuestas y guías revisadas; dudas con responsables.'],
['23 oct.','Dirección y responsables: revisar preparación y autorizar el primer piloto.','Alcance de prueba, casos, participantes y criterios de salida aprobados.'],
['26 oct.–6 nov.','Servicio piloto y coordinación: propuesta de ensayo del gestor y correcciones.','Un caso normal, dos excepciones y recorrido completo registrados. Ventana a confirmar el 29 sep.'],
['9–20 nov.','Dirección y servicios: ampliar progresivamente si el piloto resulta aceptable.','Incorporación por grupos y revisión de dependencias entre unidades.'],
['27 nov.','Dirección y responsables: consolidar aceptación funcional del alcance acordado.','Acta de aceptación o lista concreta de correcciones; cierre de nuevas funciones según plan.'],
['4 dic.','Coordinación y responsables: comprobar capacitación y preparación operativa.','Decisión documentada de iniciar o aplazar; recuperación y soporte comprobados.'],
['7–11 dic.','Equipos autorizados: uso gradual y estabilización.','Incidencias atendidas y seguimiento diario durante el arranque.'],
['14–15 dic.','Dirección: revisar resultados, pendientes y continuidad.','Aceptación formal del alcance aprobado y responsables del seguimiento.']])
para(p,'Si un requisito de salida no se cumple, se registra el motivo y se reprograma antes de ampliar el uso. La referencia de 354 horas del plan (295 base y 59 de reserva) es provisional e institucional; no corresponde a cada servicio ni sustituye la confirmación de disponibilidad real.')

p=page('Cómo sabremos si funciona','Aprendizaje del piloto y criterios de aceptación')
head(p,'7. Resultados que se propone medir')
para(p,'Se propone revisar estos indicadores cada semana y confirmarlos con Dirección antes del piloto. No son resultados ya obtenidos. La primera medición establecerá el punto de partida; el número de formularios llenos, por sí solo, no demuestra éxito.')
table(p,['Indicador','Cómo comprobarlo y meta propuesta'],[
['Responsabilidad clara','El 100 % de los procesos seleccionados tiene responsable, suplente y revisor identificados. Comprueba: coordinación.'],
['Información confiable','El 100 % de los procesos del piloto tiene descripción revisada y formatos autorizados o faltantes expresamente resueltos. Comprueba: revisor del servicio.'],
['Uso comprensible','Cada participante completa su recorrido asignado y explica qué hacer ante una duda; se registra la ayuda requerida y se repite si es necesario. Comprueba: facilitador.'],
['Compromisos realizables','Toda tarea comprometida tiene resultado, responsable, fecha y esfuerzo; se revisan conflictos de disponibilidad. Comprueba: responsable del servicio.'],
['Decisiones atendidas','Ningún bloqueo crítico carece de responsable y próxima fecha de revisión. Comprueba: coordinación con Dirección.'],
['Piloto satisfactorio','Un caso normal y dos excepciones ejecutados, documentos y permisos verificados, y saldos cuando correspondan al proceso. Sin fallas críticas abiertas. Comprueba: responsable y revisor.']])
head(p,'Condiciones para aceptar y ampliar')
para(p,'El responsable deberá registrar qué se probó, qué ocurrió y qué correcciones se hicieron. La aceptación requiere que el recorrido seleccionado pueda completarse, que los accesos sean los adecuados y que no existan fallas que comprometan la información o impidan el trabajo esencial. Se revisarán también tareas, suplencias, avisos internos e informes incluidos en el ensayo.')
para(p,'El acta indicará el alcance aceptado, participantes, fecha, resultados y pendientes menores con responsable y plazo. Un cambio posterior que afecte lo aceptado exige repetir las pruebas correspondientes. La autorización para atención real con módulos clínicos es una decisión adicional; no se obtiene al aprobar el gestor.')
head(p,'Riesgos prácticos y respuesta')
para(p,'Falta de tiempo: reservar bloques y ajustar compromisos antes del vencimiento. Respuestas contradictorias: documentar la diferencia y resolverla con el responsable. Dificultad de uso: acompañar el primer ejercicio y mejorar la orientación. Ausencia de una persona: activar la suplencia autorizada y revisar el trabajo pendiente. Crecimiento del alcance: priorizar con Dirección y reestimar, sin sumar funciones a fechas ya comprometidas por simple suposición.')
para(p,'Antes del uso operativo se requieren entorno autorizado y aislado, soporte y recuperación comprobados. Esta preparación deberá preservar las demás aplicaciones del VPS. Los avisos externos todavía requieren definir canal, autorización, implementación y pruebas; mientras tanto, la organización del piloto debe contemplar la consulta de los avisos internos.')

p=page('Acuerdos para Dirección y horizonte del proyecto','Reunión del 29 de septiembre y continuidad')
head(p,'8. Decisiones que se propone dejar por escrito')
table(p,['Decisión','Acuerdo que se necesita'],[
['Piloto y responsables','Servicio, proceso inicial, titular, suplente, revisor y participantes.'],
['Tiempo y prioridades','Disponibilidad por persona, procesos críticos y fecha del primer ejercicio.'],
['Información y apoyo','Documentos permitidos, datos del ensayo, enlace de soporte y medio de atención.'],
['Aceptación y operación','Quién autoriza cada hito, criterios del piloto y condiciones del entorno operativo.'],
['Avisos y seguimiento','Canal externo que se evaluará, si se requiere, y próxima revisión de avances.']])
para(p,'Agenda sugerida de 45 minutos: propósito y alcance (10), ejemplo de participación (10), elección del piloto y responsables (15), calendario, apoyo y acuerdos (10). Duración propuesta, pendiente de confirmación. El cierre debe producir una lista de decisiones y tareas; no sólo una demostración de pantallas.')
head(p,'9. Horizonte del sistema completo: referencia original')
para(p,'Estas son ventanas del plan original para módulos especializados. Son distintas del ensayo del gestor y están sujetas a validación y reestimación; no se presentan como funciones ya disponibles ni como compromisos nuevos aprobados.')
table(p,['Unidades / frente','Módulo o integración','Piloto original'],[
['Odontología y Fisioterapia','12–23 oct.','26 oct.–6 nov.'],
['Nutrición y Psicología','19–30 oct.','2–13 nov.'],
['Deporte, UAPS, Comunidad y Jurídico¹','26 oct.–6 nov.','9–20 nov.'],
['Administración, docencia y recursos','Núcleo: 1–9 oct.\nIntegración: 19 oct.–13 nov.','26 oct.–20 nov.']])
para(p,'¹ La inclusión de Jurídico debe confirmarse. El calendario general prevé funciones especializadas al 6 de noviembre, integraciones y migración de muestra al 13, y pilotos e integración de servicios al 20. Su viabilidad debe revisarse el 9 de octubre con información real de alcance y capacidad.')
head(p,'Fuentes y estado de este documento')
para(p,'Elaborado a partir de Plan_Trabajo_Escuela_Salud_Modelo.docx, versión 3.1 del 23 de septiembre de 2026 (incluidos C21–C27), Prompt_Maestro_Salud_Modelo_IA (1).md, el registro de decisiones y pendientes del proyecto, y la reunión comunicada por el solicitante. El cronograma de adopción, la agenda, la capacitación y las metas de medición son propuestas para acuerdo institucional.')
para(p,'Documento de presentación y trabajo, versión 1.0. No constituye acta de aceptación ni atribuye nombramientos pendientes. Su finalidad es que Dirección y los equipos comprendan el propósito, conozcan su participación y puedan acordar un inicio ordenado.')

# Editable Word
D=Document(); sec=D.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.6); sec.left_margin=sec.right_margin=Inches(.7)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
normal=D.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',25),('Heading 1',19),('Heading 2',12)]:
 s=D.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('163B50')
header=sec.header.paragraphs[0];header.text='ESCUELA DE SALUD  |  SALUD MODELO';header.style='Caption'
f=sec.footer.paragraphs[0];f.text='Propuesta institucional • 28/09/2026                                       Página '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for i,p in enumerate(pages):
 if i:D.add_page_break()
 D.add_heading(p['title'],0 if i==0 else 1);D.add_paragraph(p['subtitle'],'Subtitle')
 for kind,val in p['items']:
  if kind=='p':D.add_paragraph(val)
  elif kind=='h':D.add_heading(val,2)
  else:
   headers,rows=val;t=D.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
   for c,v in zip(t.rows[0].cells,headers):c.text=v
   repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
   for row in rows:
    cells=t.add_row().cells
    for c,v in zip(cells,row):c.text=v
   for row in t.rows:
    tr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(tr)
    for c in row.cells:
     for par in c.paragraphs:
      par.paragraph_format.space_after=Pt(4);par.paragraph_format.space_before=Pt(3)
      for r in par.runs:r.font.size=Pt(9)
 D.add_paragraph('')
D.core_properties.title='Salud Modelo: exposición de motivos y plan de participación y adopción';D.core_properties.subject='Documento para Dirección y colaboradores';D.core_properties.author='Proyecto Salud Modelo'
base='Salud_Modelo_Documento_para_Direccion_y_Colaboradores'
D.save(OUT/(base+'.docx'))
# PDF independently typeset from identical content; no system services needed.
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodySM',fontName='Helvetica',fontSize=9.3,leading=12.4,spaceAfter=7,textColor=colors.HexColor('#263640')))
styles.add(ParagraphStyle(name='TitleSM',fontName='Helvetica-Bold',fontSize=22,leading=25,spaceAfter=9,textColor=colors.HexColor('#163B50')))
styles.add(ParagraphStyle(name='SubSM',fontName='Helvetica',fontSize=11,leading=14,spaceAfter=13,textColor=colors.HexColor('#52717C')))
styles.add(ParagraphStyle(name='HeadSM',fontName='Helvetica-Bold',fontSize=11.5,leading=14,spaceBefore=6,spaceAfter=7,textColor=colors.HexColor('#163B50')))
styles.add(ParagraphStyle(name='CellSM',fontName='Helvetica',fontSize=8.3,leading=10.7))
styles.add(ParagraphStyle(name='CellHeadSM',parent=styles['CellSM'],fontName='Helvetica-Bold',textColor=colors.white))
def P(s,style):return Paragraph(escape(s).replace('\n','<br/>'),styles[style])
story=[]
for i,p in enumerate(pages):
 if i:story.append(PageBreak())
 story += [P(p['title'],'TitleSM'),P(p['subtitle'],'SubSM')]
 for kind,val in p['items']:
  if kind in ('p','h'):story.append(P(val,'BodySM' if kind=='p' else 'HeadSM'))
  else:
   headers,rows=val;n=len(headers);widths=([130,381] if n==2 else ([78,230,203] if i==5 else [153,185,173]))
   data=[[P(x,'CellHeadSM') for x in headers]]+[[P(x,'CellSM') for x in row] for row in rows]
   t=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
   t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#163B50')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#EFF4F6'),colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('LINEBELOW',(0,-1),(-1,-1),.4,colors.HexColor('#B8CAD2'))]))
   story += [t,Spacer(1,8)]
def footer(c,doc):
 c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#52717C'));c.drawString(42,815,'ESCUELA DE SALUD  |  SALUD MODELO');c.drawString(42,26,'Propuesta institucional • 28/09/2026');c.drawRightString(553,26,f'Página {doc.page}')
SimpleDocTemplate(str(OUT/(base+'.pdf')),pagesize=(595.28,841.89),rightMargin=42,leftMargin=42,topMargin=45,bottomMargin=43,title=D.core_properties.title,author='Proyecto Salud Modelo').build(story,onFirstPage=footer,onLaterPages=footer)
# Readable text companion for maintenance and review.
lines=[]
for p in pages:
 lines += ['# '+p['title'],p['subtitle'],'']
 for k,v in p['items']:
  if k=='t':
   h,rows=v;lines+=[' | '.join(h)]+[' | '.join(r).replace('\n',' ') for r in rows]+['']
  else:lines += [('## ' if k=='h' else '')+v,'']
(OUT/(base+'.md')).write_text('\n'.join(lines))
print('CREATED',base)
print('CONTENT',len(' '.join(lines).split()),'words;',len(pages),'sections')
